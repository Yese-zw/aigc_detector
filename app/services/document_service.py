"""Document parsing and quota calculation."""

import io
import re

from docx import Document
from fastapi import HTTPException, status

from app.core.logger import logger


class DocumentService:
    def extract_text(self, file_content: bytes, filename: str) -> str:
        lower_name = filename.lower()
        if lower_name.endswith(".txt"):
            return self._extract_txt(file_content)
        if lower_name.endswith(".docx"):
            return self._extract_docx(file_content)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="仅支持txt或docx格式的文件")

    def count_billable_chars(self, text: str) -> int:
        return len(re.findall(r"[\u4e00-\u9fa5a-zA-Z0-9]", text))

    def upload_quota_cost(self, text: str, mode: str) -> int:
        cost = self.count_billable_chars(text)
        return cost * 2 if mode == "3" or mode == 3 else cost

    def _extract_txt(self, file_content: bytes) -> str:
        try:
            return file_content.decode("utf-8")
        except UnicodeDecodeError:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="TXT文件编码错误，仅支持UTF-8编码")

    def _extract_docx(self, file_content: bytes) -> str:
        try:
            doc = Document(io.BytesIO(file_content))
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += "\n" + cell.text
            return text
        except Exception as exc:
            logger.error(f"解析docx文件失败: {str(exc)}")
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="docx文件解析失败，可能是损坏的文件或非docx格式")
